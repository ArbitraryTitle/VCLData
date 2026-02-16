#!/usr/bin/env python3
from pptx import Presentation
import json

prs = Presentation('64-2025-05-22 VCL Meeting.pptx')

# extract tables from slides
tables_data = []

for slide_num, slide in enumerate(prs.slides, 1):
    for shape in slide.shapes:
        if shape.has_table:
            table = shape.table
            rows = []
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    cells.append(cell.text.strip())
                rows.append(cells)

            tables_data.append({
                'slide': slide_num,
                'rows': rows
            })

# print all tables
for table_info in tables_data:
    print(f"\n=== Slide {table_info['slide']} ===")
    for row in table_info['rows']:
        print(row)

# save as json
with open('tables_raw.json', 'w') as f:
    json.dump(tables_data, f, indent=2)
