#!/usr/bin/env python3
from pptx import Presentation
import json

prs = Presentation('64-2025-05-22 VCL Meeting.pptx')

# extract all text from slides
data = []
for i, slide in enumerate(prs.slides):
    slide_data = {
        'slide_num': i + 1,
        'text': []
    }
    for shape in slide.shapes:
        if hasattr(shape, "text"):
            if shape.text.strip():
                slide_data['text'].append(shape.text.strip())
    if slide_data['text']:
        data.append(slide_data)

# print raw data first
print(json.dumps(data, indent=2))
